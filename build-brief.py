import re
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = Path(__file__).parent
doc = Document()
sec = doc.sections[0]
sec.top_margin = sec.bottom_margin = Inches(.7)
sec.left_margin = sec.right_margin = Inches(.8)
for name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'List Bullet', 'List Number']:
    s = doc.styles[name]
    s.font.name = 'Calibri'
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.font.size = Pt(11)
    s.paragraph_format.space_after = Pt(6)
    s.paragraph_format.line_spacing = 1.08
doc.styles['Title'].font.size = Pt(25)
doc.styles['Heading 1'].font.size = Pt(16)
doc.styles['Heading 2'].font.size = Pt(12)
for name in ['Heading 1', 'Heading 2']:
    doc.styles[name].paragraph_format.space_before = Pt(12)
    doc.styles[name].paragraph_format.keep_with_next = True

def runs(p, text):
    for i, chunk in enumerate(re.split(r'\*\*(.*?)\*\*', text)):
        r = p.add_run(chunk)
        r.bold = i % 2 == 1

for line in (root/'EONLINK-Developer-Brief.md').read_text(encoding='utf-8').splitlines():
    if not line.strip():
        continue
    if line.startswith('# '):
        doc.add_paragraph('EONLINK', 'Title')
        doc.add_paragraph('Product and Development Scope Brief', 'Heading 1')
        continue
    if line.startswith('## '):
        text = re.sub(r'^\d+\.\s*', '', line[3:]).replace(' — ', ' ').replace('/', ' and ')
        doc.add_paragraph(text, 'Heading 1')
    elif line.startswith('### '):
        text = re.sub(r'^[A-F]\.\s*', '', line[4:]).replace(' — ', ' ').replace(',', '')
        doc.add_paragraph(text, 'Heading 2')
    elif line.startswith('- '):
        runs(doc.add_paragraph(style='List Bullet'), line[2:])
    elif re.match(r'^\d+\. ', line):
        runs(doc.add_paragraph(style='List Number'), re.sub(r'^\d+\. ', '', line))
    else:
        runs(doc.add_paragraph(), line)

p = sec.footer.paragraphs[0]
p.alignment = 2
r = p.add_run()
f = OxmlElement('w:fldSimple'); f.set(qn('w:instr'), 'PAGE')
r._r.addnext(f)
doc.core_properties.title = 'EONLINK Product and Development Scope Brief'
doc.core_properties.author = 'EONLINK'
doc.save(root/'EONLINK-Developer-Brief.docx')
print(root/'EONLINK-Developer-Brief.docx')
